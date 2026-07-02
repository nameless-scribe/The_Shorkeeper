#!/usr/bin/env python3
"""Generate The Shorekeeper user guide as a Word-compatible .docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUTPUT = Path(__file__).resolve().parent.parent / "docs" / "守岸人-使用说明.docx"


def set_default_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbers(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, text in enumerate(row):
            cells[c_idx].text = text


def add_link(doc: Document, text: str, url: str) -> None:
    p = doc.add_paragraph()
    part = p.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = docx_hyperlink(r_id, text)
    p._p.append(hyperlink)


def docx_hyperlink(r_id: str, text: str):
    from docx.oxml import OxmlElement

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    return hyperlink


def build() -> Document:
    doc = Document()
    set_default_font(doc)

    title = doc.add_heading("The Shorekeeper 用户使用说明", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("应用名称：The Shorekeeper（守岸人）\n文档版本：1.0  |  更新日期：2026-07-02  |  应用版本：v0.1.0")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_heading("1. 产品简介", level=1)
    doc.add_paragraph(
        "The Shorekeeper（守岸人）是一款自用桌面 AI Agent 应用。它支持流式对话、工具调用、长期记忆与知识库检索（RAG）、"
        "联网搜索、定时任务、文档生成、好感度阶段、多窗口伴侣界面等功能。"
    )
    doc.add_paragraph("核心能力包括：")
    add_bullets(
        doc,
        [
            "长期记忆、Worldbook 触发词背景、用户画像",
            "泰提斯终端知识库（MD / TXT / PDF / DOCX 导入与混合检索）",
            "博查联网搜索、网页抓取、本地文件读写、多格式文档导出",
            "Agent Skills 技能包与 MCP 外部工具扩展",
            "CosyVoice 语音朗读（可选）",
            "主题外观自定义、Token 统计、定时提醒与静默 Agent 任务",
        ],
    )

    doc.add_heading("2. 安装与首次使用", level=1)
    doc.add_heading("2.1 环境要求（开发者 / 源码运行）", level=2)
    add_bullets(doc, ["Node.js 20 及以上", "pnpm 包管理器", "Windows 系统（当前主要支持平台）"])

    doc.add_heading("2.2 安装包方式（推荐普通用户）", level=2)
    add_numbers(
        doc,
        [
            "获取 release/ 目录下的 Windows 安装包（.exe）",
            "双击安装，无需安装 Node.js 或 pnpm",
            "首次启动后自动创建数据目录（默认 D:\\SQLlite\\）",
            "打开 设置 → API 设置，填写模型 API Key 与接入地址",
            "（可选）设置 → 人设，编辑角色与 System Prompt",
            "（可选）设置 → 外观，切换主题或上传背景/头像",
            "开始与守岸人对话",
        ],
    )

    doc.add_heading("2.3 源码开发方式", level=2)
    add_numbers(
        doc,
        [
            "复制 .env.example 为 .env，填入百炼 API Key 与 Base URL",
            "运行 pnpm install 安装依赖",
            "运行 pnpm db:init 初始化数据库",
            "运行 pnpm dev 启动开发模式",
        ],
    )

    doc.add_heading("2.4 首次使用建议配置顺序", level=2)
    add_numbers(
        doc,
        [
            "设置 → API 设置：填写百炼 Key、Base URL、模型 ID，确认顶栏显示「已连接」",
            "设置 → 人设：确认或调整角色名与 System Prompt",
            "设置 → 用户信息：填写称呼、偏好等画像字段",
            "设置 → 插件：按需开启联网搜索、本地文件等工具",
            "设置 → 技能：启用需要的 Agent Skills",
            "设置 → 泰提斯终端：导入知识库文档（可选）",
            "设置 → 语音：配置 CosyVoice 朗读（可选）",
            "在主聊天窗输入消息，Enter 发送",
        ],
    )

    doc.add_heading("3. 应用窗口与导航", level=1)
    add_table(
        doc,
        ["窗口", "用途", "如何打开"],
        [
            ["聊天窗", "主对话界面", "启动默认显示；Dock 点头像；托盘「显示聊天窗」；双击托盘图标"],
            ["状态面板", "角色状态、喂食、Token 用量", "Dock 点「状态」条；托盘「显示状态面板」"],
            ["日程面板", "Token 统计图表、定时任务一览", "Dock 点「日程/Token」条；托盘「显示日程面板」"],
            ["Dock 悬浮条", "快捷入口与迷你信息", "聊天/状态/日程全部隐藏时自动出现"],
            ["提醒窗", "定时提醒通知", "定时任务到点自动弹出"],
        ],
    )
    doc.add_paragraph(
        "重要：点击任意主窗口的关闭按钮（×）不会退出程序，而是隐藏到系统托盘。"
        "要彻底退出，请在托盘右键菜单选择「退出」。"
    )

    doc.add_page_break()

    doc.add_heading("4. 主聊天界面", level=1)
    doc.add_heading("4.1 顶栏（TitleBar）", level=2)
    add_table(
        doc,
        ["元素", "功能"],
        [
            ["标题 The Shorekeeper", "应用名称"],
            ["副标题", "已连接时显示配置名与模型名；未配置时显示「未配置 API Key」"],
            ["☰", "切换历史会话侧栏"],
            ["＋", "新建对话（运行中会先中止当前任务）"],
            ["⚙", "打开设置中心"],
            ["─", "最小化窗口"],
            ["×", "关闭窗口（隐藏到托盘，不退出程序）"],
        ],
    )

    doc.add_heading("4.2 工作流状态条", level=2)
    doc.add_paragraph(
        "Agent 运行时显示当前阶段：准备 → 思考 → 工具 → 输出。空闲时显示「就绪 · 等待你的消息」。"
        "若工具需要用户批准（如写入文件），会弹出权限确认对话框，可选择「允许」或「拒绝」。"
    )

    doc.add_heading("4.3 消息列表", level=2)
    add_bullets(
        doc,
        [
            "用户消息显示在右侧，带用户头像；可附带文件卡片",
            "助手消息显示在左侧，带 Agent 头像；可显示工具调用卡片与相关文件",
            "流式输出时有光标动画；思考中显示「思考中」动画",
            "助手回复完成后，可点击朗读按钮朗读该条消息（需先在语音设置中启用）",
            "开启自动朗读时，每轮回复结束会自动播放",
        ],
    )

    doc.add_heading("4.4 输入栏", level=2)
    add_table(
        doc,
        ["操作", "方法"],
        [
            ["发送消息", "点击「发送」或按 Enter（Shift+Enter 换行）"],
            ["上传附件", "点击附件按钮或拖入文件到输入区"],
            ["切换模型", "输入栏左侧模型下拉（多套 API 配置时显示）"],
            ["移除附件", "点击附件标签上的 ×"],
        ],
    )
    doc.add_paragraph(
        "支持的附件类型包括：txt、md、json、csv、doc/docx、xls/xlsx 等；单文件上限约 20MB。"
        "Agent 运行中或未配置 API Key 时，输入栏不可用。"
    )

    doc.add_heading("4.5 历史会话侧栏", level=2)
    add_bullets(
        doc,
        [
            "搜索会话标题或消息内容",
            "切换、归档、删除历史会话",
            "清理无消息的空会话",
            "显示「已压缩」「已归档」等标记",
        ],
    )

    doc.add_page_break()

    doc.add_heading("5. 设置中心", level=1)
    doc.add_heading("5.1 打开方式", level=2)
    add_bullets(
        doc,
        [
            "聊天窗顶栏设置按钮",
            "状态面板底部「设置」链接",
            "日程面板「任务设置 →」链接（直接跳转到定时任务页）",
        ],
    )
    doc.add_paragraph("设置以全屏抽屉形式打开，左侧为导航分组，右侧为具体设置内容。")

    doc.add_heading("5.2 设置页一览", level=2)
    add_table(
        doc,
        ["分组", "页面", "说明"],
        [
            ["能力", "插件", "联网搜索、本地文件、文档生成等工具开关"],
            ["能力", "技能", "Agent Skills 技能包启用/禁用"],
            ["能力", "MCP", "Model Context Protocol 外部服务器"],
            ["人格与记忆", "人设", "角色名与 System Prompt（每轮注入）"],
            ["人格与记忆", "用户信息", "长期用户画像字段"],
            ["人格与记忆", "记忆", "RAG 模式、自动记忆提取、上下文压缩"],
            ["人格与记忆", "Worldbook", "触发词与背景设定"],
            ["个性化", "外观", "9 套主题、壁纸、头像、遮罩"],
            ["个性化", "语音", "CosyVoice 朗读与复刻音色"],
            ["数据与任务", "泰提斯终端", "知识库导入与向量配置"],
            ["数据与任务", "定时任务", "提醒通知与静默 Agent 任务"],
            ["系统", "API 设置", "模型协议、Base URL、API Key"],
            ["系统", "免责声明", "使用须知"],
        ],
    )

    doc.add_page_break()

    doc.add_heading("6. 各设置页详细说明", level=1)

    doc.add_heading("6.1 API 设置", level=2)
    doc.add_paragraph("配置对话模型接入。支持保存多套配置并随时切换。")
    add_numbers(
        doc,
        [
            "填写配置名称、Base URL、模型 ID（小写）、API Key",
            "选择请求协议：OpenAI 兼容（百炼/DeepSeek 等）或 Anthropic（Claude 风格）",
            "点击「保存配置」",
            "在配置卡片列表中点击即可切换当前使用的模型",
        ],
    )
    doc.add_paragraph("应用内配置优先于 .env 文件。顶栏显示「已连接」表示配置成功。")

    doc.add_heading("6.2 人设", level=2)
    doc.add_paragraph(
        "编辑角色名与 System Prompt。人设决定 Agent 每一轮如何扮演，与 Worldbook 触发背景互补，不可替代。"
    )
    add_bullets(
        doc,
        [
            "保存后下一条消息起生效",
            "可「恢复默认」为内置守岸人人设",
            "System Prompt 上限 16,000 字",
        ],
    )

    doc.add_heading("6.3 用户信息", level=2)
    doc.add_paragraph("管理长期用户画像（键值对），如称呼、偏好语气、简介等。画像会注入每次对话的 system prompt。")

    doc.add_heading("6.4 记忆（RAG 与上下文）", level=2)
    add_table(
        doc,
        ["设置项", "说明", "默认值"],
        [
            ["启用 RAG 检索注入", "知识库检索总开关", "开"],
            ["RAG 注入模式", "目录（省 Token）/ 自动 / 工具", "目录"],
            ["RAG 最低相似度", "0.2–0.8，越高越严格", "0.35"],
            ["自动提取记忆", "每轮 / 每 N 轮 / 手动", "每 N 轮"],
            ["间隔 N 轮", "自动提取间隔", "3"],
            ["最近消息条数", "送入模型的历史消息数", "20"],
            ["压缩阈值", "超过后压缩旧消息为摘要", "30"],
        ],
    )
    doc.add_paragraph("「目录」模式仅注入文档文件名列表，模型按需调用 search_knowledge 工具，Token 开销最低。")

    doc.add_heading("6.5 Worldbook", level=2)
    doc.add_paragraph(
        "当用户消息命中触发词时，自动注入对应背景设定。可添加、编辑、启用/禁用、删除条目。"
        "触发词用逗号分隔，如：魔法,法术,咒语。"
    )

    doc.add_heading("6.6 外观", level=2)
    doc.add_paragraph(
        "9 套主题预设：守岸人·星空（默认）、午夜紫、拂晓、深海青、晨雾绿、桜夜、薄暮紫、板岩灰、霓虹脉冲。"
    )
    add_bullets(
        doc,
        [
            "可上传自定义背景壁纸（拖拽或从本机选择）",
            "适应方式：铺满窗口 / 完整显示",
            "遮罩强度：0–100%，保证聊天气泡可读",
            "分别更换 Agent 与用户头像（Agent 头像同步到 Dock 与托盘）",
        ],
    )
    doc.add_paragraph("自定义资源保存在 appearance/ 目录，与数据库同根目录。")

    doc.add_heading("6.7 语音", level=2)
    doc.add_paragraph("对话模型与语音合成相互独立：可用 Claude 等模型聊天，同时用百炼 CosyVoice 朗读。")
    add_numbers(
        doc,
        [
            "选择「独立百炼（推荐）」或「复用对话 API」",
            "填写百炼 API Key 与 CosyVoice 合成接入点",
            "开启「启用语音朗读」",
            "粘贴百炼声音复刻 voice_id 并保存",
            "点击「试听示例句」验证（「调律者，我在这里。」）",
            "可选：开启「自动朗读」",
        ],
    )
    doc.add_paragraph("合成模型须与百炼复刻时的 target_model 一致（默认 cosyvoice-v3.5-plus）。")

    doc.add_heading("6.8 泰提斯终端（知识库）", level=2)
    add_numbers(
        doc,
        [
            "配置向量 API（Embedding）：可与对话相同，或单独配置百炼 Embedding",
            "点击「测试连接」确认向量 API 可用",
            "选择 MD / TXT / DOCX / PDF 文件导入",
            "在「记忆」页选择合适的 RAG 注入模式",
            "对话中提问与文档相关的内容",
        ],
    )
    doc.add_paragraph(
        "若对话使用 Claude 等第三方代理，通常需单独配置百炼 Embedding。"
        "导入后若更换 Embedding 模型，需「重建全部向量」。"
    )

    doc.add_heading("6.9 定时任务", level=2)
    add_table(
        doc,
        ["任务类型", "行为"],
        [
            ["提醒通知", "到点弹出提醒窗，显示标题与正文"],
            ["Agent 静默执行", "后台执行 Agent 提示词，不打开聊天窗"],
        ],
    )
    doc.add_paragraph(
        "支持周期性（Cron 表达式，如 0 9 * * * = 每天 9:00）与仅一次任务。"
        "Agent 也可在对话中通过工具创建定时任务。"
    )

    doc.add_heading("6.10 插件", level=2)
    add_table(
        doc,
        ["插件", "功能", "配置要点"],
        [
            ["联网搜索", "博查 API 联网检索", "填写博查 API Key（open.bochaai.com）"],
            ["本地文件", "读写 Agent 工作区文件", "只读 / 审批（需确认）/ 完全（免确认）"],
            ["多格式编写", "导出 Excel/Word/PDF/Markdown", "开关即可"],
            ["网页抓取", "fetch_url 转文本", "开关即可"],
            ["收支记账", "bookkeeping 工具", "开关即可"],
            ["生活工具", "旅行规划、天气、翻译", "可展开查看子工具"],
            ["MCP 工具", "已注册 MCP Server 的工具", "在 MCP 页管理服务器"],
        ],
    )

    doc.add_heading("6.11 技能", level=2)
    doc.add_paragraph(
        "技能来自内置 skills/ 目录（安装包已打包）。启用后注入 system prompt。"
        "多个技能同时启用时，可用工具为各白名单的并集。"
    )

    doc.add_heading("6.12 MCP", level=2)
    doc.add_paragraph(
        "添加 MCP Server：填写名称、启动命令、参数 JSON、环境变量 JSON。"
        "支持启用/禁用、测试连接、删除。工具名格式为 mcp__服务器__工具。"
    )

    doc.add_heading("6.13 免责声明", level=2)
    doc.add_paragraph(
        "本软件为个人自用桌面 AI Agent，仅供学习与私人辅助。"
        "Agent 生成内容、工具执行结果（文件写入、网络请求等）由用户自行承担风险。"
        "联网搜索、网页抓取、MCP 第三方工具可能访问外部服务，请注意 API Key 与数据安全。"
    )

    doc.add_page_break()

    doc.add_heading("7. Dock 悬浮条与辅助面板", level=1)
    doc.add_heading("7.1 Dock 悬浮条", level=2)
    doc.add_paragraph("当聊天窗、状态面板、日程面板全部隐藏时，Dock 自动出现在屏幕上。")
    add_table(
        doc,
        ["操作", "方法"],
        [
            ["打开聊天窗", "点击 Agent 头像"],
            ["打开状态面板", "点击「状态」信息条"],
            ["打开日程面板", "点击「日程」或「今日 Token」信息条"],
            ["移动 Dock", "拖动空白区域（未固定位置时）"],
            ["窗口置顶", "悬停后点击置顶按钮"],
            ["固定位置", "悬停后点击固定按钮，锁定后不可拖动"],
        ],
    )

    doc.add_heading("7.2 状态面板", level=2)
    doc.add_paragraph("显示角色头像、在线/离线状态、当前模型、状态/心情/羁绊卡片、今日 Token 用量。")
    add_bullets(
        doc,
        [
            "喂食：触发「进食中」状态",
            "可快速打开聊天或设置",
            "角色名来自 设置 → 人设",
        ],
    )

    doc.add_heading("7.3 日程面板", level=2)
    add_bullets(
        doc,
        [
            "今日 Token 用量、本周/累计统计、日预算进度",
            "近 7 日 Token 折线图",
            "定时任务列表（名称、调度、启用状态）",
            "「任务设置 →」跳转至设置中的定时任务页",
        ],
    )

    doc.add_page_break()

    doc.add_heading("8. 系统托盘", level=1)
    doc.add_paragraph("应用隐藏到托盘后，可通过托盘图标快速恢复窗口。")
    add_table(
        doc,
        ["操作", "方法"],
        [
            ["显示聊天窗", "双击托盘图标，或右键 →「显示聊天窗」"],
            ["显示状态面板", "右键 →「显示状态面板」"],
            ["显示日程面板", "右键 →「显示日程面板」"],
            ["退出应用", "右键 →「退出」"],
        ],
    )
    doc.add_paragraph("托盘图标优先使用 Agent 自定义头像。")

    doc.add_heading("9. 数据目录与文件路径", level=1)
    add_table(
        doc,
        ["路径", "说明"],
        [
            ["D:\\SQLlite\\", "数据库根目录（可用 SHOREKEEPER_DB_DIR 修改）"],
            ["D:\\SQLlite\\shorekeeper.db", "聊天、设置、记忆、会话等数据"],
            ["D:\\SQLlite\\workspace\\", "Agent 读写与生成文件的工作区"],
            ["D:\\SQLlite\\appearance\\", "自定义背景、头像、语音样本等资源"],
        ],
    )
    doc.add_paragraph("若目标机没有 D: 盘，会自动回退到用户目录下的应用数据文件夹。备份时建议拷贝整个数据目录。")

    doc.add_heading("10. 环境变量参考（高级）", level=1)
    doc.add_paragraph("应用内设置优先于 .env。修改 .env 后通常需重启应用。")
    add_table(
        doc,
        ["变量", "说明"],
        [
            ["SHOREKEEPER_DB_DIR / SHOREKEEPER_DB_PATH", "数据库目录与文件路径"],
            ["SHOREKEEPER_WORKSPACE_DIR", "Agent 工作区路径"],
            ["OPENAI_API_KEY / OPENAI_BASE_URL", "对话 API Key 与接入点"],
            ["DEFAULT_MODEL", "默认模型 ID（如 qwen3.6-plus）"],
            ["EMBEDDING_API_KEY / EMBEDDING_BASE_URL", "向量 API（RAG 用，可与对话分离）"],
            ["RAG_ENABLED / MAX_HISTORY_MESSAGES 等", "RAG 与记忆相关默认值"],
            ["WEB_SEARCH_API_KEY / BOCHA_API_KEY", "博查联网搜索 Key"],
            ["VOICE_API_KEY / VOICE_TTS_ENDPOINT", "CosyVoice 语音合成配置"],
        ],
    )

    doc.add_page_break()

    doc.add_heading("11. 常见问题", level=1)
    add_table(
        doc,
        ["现象", "解决方法"],
        [
            ["无法发送消息", "检查 设置 → API 设置；确认顶栏非「未配置 API Key」"],
            ["401 / 404 模型错误", "检查 Key 完整性；模型 ID 使用小写 API 名称"],
            ["关闭窗口后找不到应用", "查看系统托盘；双击或右键「显示聊天窗」"],
            ["所有窗口都关了", "Dock 悬浮条会出现，点击头像恢复聊天"],
            ["知识库检索无效", "确认已导入文档；检查 Embedding 配置；必要时重建向量"],
            ["Claude 对话 + RAG 失败", "泰提斯终端 → 向量 API 选「单独配置」百炼 Embedding"],
            ["文件写入被拦截", "设置 → 插件 → 本地文件改为「审批」或「完全」"],
            ["语音无朗读按钮", "设置 → 语音 → 启用朗读 + 配置 Key 与 voice_id"],
            ["语音音量偏小", "合成音量拉至 100，适当提高播放增益"],
            ["缓存始终未命中", "部分 API 不返回 cached_tokens，属正常提示"],
        ],
    )

    doc.add_heading("12. 相关链接", level=1)
    add_link(doc, "阿里云百炼控制台", "https://dashscope.console.aliyun.com/")
    add_link(doc, "博查开放平台（联网搜索）", "https://open.bochaai.com")
    add_link(
        doc,
        "CosyVoice 声音复刻文档",
        "https://help.aliyun.com/zh/model-studio/cosyvoice-clone-design-api",
    )

    doc.add_paragraph("")
    end = doc.add_paragraph("— 文档结束 —")
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in end.runs:
        run.italic = True

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build()
    doc.save(OUTPUT)
    print(f"Generated: {OUTPUT} ({OUTPUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
